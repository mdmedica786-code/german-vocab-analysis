#!/usr/bin/env python3
"""bank_tools - build tool for Medica's German banks (vocabulary, phrases, Redemittel).

Commands: status [--write] | next <bank> <n> [--from A --to B | --wave W | --cat ID] [--incomplete]
          apply <patch.json> | validate [--final] | export
Banks: vocab | phrase | redemittel.   Root = parent of this file's folder (override: --root DIR or BANK_ROOT).
Special patch value: "__CLEAR__" empties a cell (empty string "" leaves a cell unchanged).
A patch row for a vocab/phrase id marks that row as reviewed even if no cell changes.
"""
import csv, json, os, re, shutil, sys, tempfile, datetime
from pathlib import Path

csv.field_size_limit(10**7)
args = sys.argv[1:]
ROOT = Path(os.environ.get("BANK_ROOT", Path(__file__).resolve().parent.parent))
if "--root" in args:
    i = args.index("--root"); ROOT = Path(args[i + 1]); del args[i:i + 2]
IN, OUT, WORK = ROOT / "input", ROOT / "bank_final", ROOT / "work"
README = IN / "README_input.txt"
FILES = {"vocab": "vocabulary_bank.csv", "phrase": "phrase_bank.csv", "redemittel": "redemittel_bank.csv"}
INPUTS = {"vocab": "vocabulary_input.csv", "phrase": "phrases_input.csv"}
RULE_KEY = {"vocab": "vocab", "phrase": "phrase", "redemittel": "redemittel"}
STATE = WORK / "state.json"
CLEAR = "__CLEAR__"

RED_HEADER = "id,wave,cat_id,category,function,speaker,audience,redemittel_de,meaning_en,level,register,example_de,example_en,grammar_note,variants,exam_use,priority,source,check".split(",")

LOCKED = {
    "vocab": {"id", "study_order", "batch", "lane", "order_freq_first", "skip", "skip_reason", "kind", "key_word",
              "nw_ref", "nw_chapter", "nw_entry", "goethe_level", "freq_rank", "freq_basis", "todo"},
    "phrase": {"id", "order_hint", "words", "source", "level_hint", "goethe_level", "corpus_freq", "quality",
               "rater_note", "nw_ref", "old_bank_id", "todo"},
    "redemittel": {"id", "wave", "cat_id"},
}
POS = "noun verb adjective adverb preposition conjunction pronoun determiner numeral particle interjection affix".split() + ["proper noun"]
LEVEL = ["A1", "A2", "B1", "B2", "C1"]
REGISTER = "neutral formal informal colloquial technical dated regional vulgar".split()
TOPICS = """T01 Person, Familie & Beziehungen|T02 Wohnen|T03 Arbeit & Beruf|T04 Ausbildung, Schule & Studium|T05 Gesundheit & Körper|T06 Essen & Trinken|T07 Einkaufen, Kleidung & Geld|T08 Reisen, Verkehr & Orte|T09 Behörden, Post & Dienstleistungen|T10 Freizeit, Sport, Kultur & Medien|T11 Natur, Umwelt, Wetter & Tiere|T12 Gesellschaft, Politik, Recht & Geschichte|T13 Zeit, Zahlen, Maße & Mengen|T14 Gefühle, Charakter & Meinungen|T15 Sprache, Kommunikation & Denken|T16 Allgemein (Funktionswörter, abstrakter Wortschatz)|M1 Medizin: Anatomie & Körpersysteme|M2 Medizin: Symptome & Krankheiten|M3 Medizin: Untersuchung, Diagnostik & Therapie|M4 Medizin: Klinik, Pflege & Recht""".split("|")
TOPIC_BY_CODE = {t.split(" ", 1)[0]: t for t in TOPICS}
FUNCTIONS = """Meinung äußern (give an opinion)|Zustimmen / Widersprechen (agree, disagree)|Gründe nennen (cause, reason)|Folgen (consequence)|Gegensatz / Einräumung (contrast, concession)|Bedingung / Zweck (condition, purpose)|Vergleichen (comparison)|Sicherheit / Zweifel (certainty, doubt)|Vorschlagen / Raten (suggest, advise)|Höflich bitten / fragen (polite request)|Hilfe anbieten / Service (offer help)|Entschuldigen / Bedauern (apology, regret)|Gefühle ausdrücken (feelings)|Probleme / Beschwerden beschreiben (problem, symptom, complaint)|Berichten (reporting)|Gliedern / Verknüpfen (structuring, discourse markers)|Zeit / Reihenfolge (time, sequence)|Ort / Richtung (location, direction)|Menge / Häufigkeit (quantity, frequency)|Können / Müssen / Dürfen (modal frames)|Begrüßen / Verabschieden / Danken (greetings, thanks)|Feste Verbindung / Redewendung (fixed expression, idiom, verb phrase)|Sonstige nützliche Wendung (other useful expression)""".split("|")
FUNC_BY_HEAD = {f.split(" (")[0]: f for f in FUNCTIONS}
SPEAKER = ["Arzt", "Patient", "Pflege", "Beide"]
AUDIENCE = ["Patient", "Angehörige", "Kollegen", "Prüfer/Publikum", "Freunde/Bekannte", "Behörde/Firma", "Alle"]
ENUMS = {
    "vocab": {"pos": POS, "level": LEVEL, "register": REGISTER, "level_basis": ["netzwerk", "goethe", "estimated"],
              "example_source": ["netzwerk", "goethe", "claude"], "medical": ["core", "related"], "topic": TOPICS},
    "phrase": {"pattern": ["fixed", "frame"], "function": FUNCTIONS, "register": REGISTER, "level": LEVEL},
    "redemittel": {"wave": ["1", "2", "3"], "function_unused": [], "speaker": SPEAKER, "audience": AUDIENCE,
                   "level": LEVEL[1:], "register": REGISTER, "priority": ["core", "extra"], "source": ["claude"]},
}
FLAG_RE = re.compile(r"^(\?[a-z]+|[a-z-]+\?|dup:[A-Za-z]\d+)$")
HTML_RE = re.compile(r"<\s*/?\s*[A-Za-z!][^>]*>|\*\*|`|^\s*#{1,6}\s|&[a-z]+;|\[[^\]]*\]\([^)]*\)")
ORDER_COLS = {
    "vocab": "id lemma kind pos article plural pres_3sg preterite perfect sep_prefix reflexive gov extra_forms meaning_en other_senses note register level level_basis topic medical key_word skip example_de example_en collocations freq_rank check todo".split(),
    "phrase": "id phrase_de pattern function register level_hint meaning_en example_de example_en grammar_note variants rater_note check todo".split(),
}


# ---------------------------------------------------------------- io
def read_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        rd = csv.reader(f)
        header = next(rd)
        rows = [dict(zip(header, r + [""] * (len(header) - len(r)))) for r in rd]
    return header, rows


def write_csv_atomic(path, header, rows):
    path = Path(path)
    fd, tmp = tempfile.mkstemp(dir=path.parent, suffix=".tmp")
    with os.fdopen(fd, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(header)
        for r in rows:
            w.writerow([r.get(h, "") for h in header])
    os.replace(tmp, path)


def load_state():
    if STATE.exists():
        return json.loads(STATE.read_text(encoding="utf-8"))
    return {"patches": 0, "reviewed": {"vocab": [], "phrase": []}, "backup_at": 0}


def save_state(s):
    STATE.write_text(json.dumps(s, ensure_ascii=False), encoding="utf-8")


def load_plan():
    _, rows = read_csv(IN / "redemittel_plan.csv")
    return rows


def load_bank(bank):
    p = OUT / FILES[bank]
    if not p.exists():
        raise SystemExit(f"missing {p} - run setup first")
    return read_csv(p)


def init_bank_final():
    OUT.mkdir(exist_ok=True); WORK.mkdir(exist_ok=True)
    for b, src in INPUTS.items():
        dst = OUT / FILES[b]
        if not dst.exists():
            shutil.copyfile(IN / src, dst)
    r = OUT / FILES["redemittel"]
    if not r.exists():
        write_csv_atomic(r, RED_HEADER, [])
    if not (OUT / "changes.csv").exists():
        with open(OUT / "changes.csv", "w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f, lineterminator="\r\n").writerow(["bank", "id", "column", "old", "new"])
    for b in INPUTS:  # recompute todo once
        h, rows = load_bank(b)
        for r in rows: r["todo"] = todo(b, r)
        write_csv_atomic(OUT / FILES[b], h, rows)


# ---------------------------------------------------------------- todo logic
def todo(bank, r):
    g = lambda k: (r.get(k) or "").strip()
    miss = []
    if bank == "vocab":
        for c in ("pos", "meaning_en", "level", "register", "topic"):
            if not g(c): miss.append(c)
        if g("level") and not g("level_basis"): miss.append("level_basis")
        kind, pos = g("kind"), g("pos")
        if kind == "noun" or (pos == "noun" and kind != "noun_noart"):
            for c in ("article", "plural"):
                if not g(c): miss.append(c)
        if pos == "verb":
            for c in ("pres_3sg", "preterite", "perfect"):
                if not g(c): miss.append(c)
        if pos == "preposition" and not g("gov"): miss.append("gov")
        if not g("skip"):
            for c in ("example_de", "example_en"):
                if not g(c): miss.append(c)
            if (g("key_word") == "yes" or g("medical") == "core") and pos in ("noun", "verb", "adjective", "adverb") and not g("collocations"):
                miss.append("collocations")
    elif bank == "phrase":
        for c in ("meaning_en", "function", "level", "register", "example_de", "example_en"):
            if not g(c): miss.append(c)
    else:
        for c in RED_HEADER:
            if c not in ("variants", "check", "id") and not g(c): miss.append(c)
    return "|".join(miss)


# ---------------------------------------------------------------- cell checks
def norm_enum(bank, col, val):
    """expand shorthand (topic code, function head); return normalised value."""
    if bank == "vocab" and col == "topic":
        return TOPIC_BY_CODE.get(val, val)
    if bank == "phrase" and col == "function":
        return FUNC_BY_HEAD.get(val.split(" (")[0], val)
    return val


def cell_errors(bank, col, val, rid):
    e = []
    if val == "": return e
    if "\t" in val or "\n" in val or "\r" in val: e.append("tab/line break")
    if HTML_RE.search(val): e.append("HTML/Markdown")
    if val != val.strip(): e.append("leading/trailing space")
    if "  " in val: e.append("double space")
    enum = ENUMS[bank].get(col)
    if enum is not None and val not in enum: e.append(f"value not allowed: {val!r}")
    if col == "check":
        for t in val.split(" "):
            if not FLAG_RE.match(t): e.append(f"bad flag {t!r}")
    if bank == "vocab":
        if col == "article" and not re.fullmatch(r"(der|die|das)(/(der|die|das))?", val): e.append("article must be der/die/das")
        if col == "meaning_en" and len(val.split(";")) > 3: e.append("meaning_en: more than 3 senses")
        if col == "note" and len(val) > 120: e.append("note > 120 chars")
        if col == "collocations" and not 1 <= len(val.split(" | ")) <= 6: e.append("collocations: use ' | '")
    if bank == "phrase":
        if col == "phrase_de" and len(val.split()) > 8: e.append("phrase_de > 8 words")
        if col == "meaning_en" and len(val.split()) > 12: e.append("meaning_en > 12 words")
        if col == "grammar_note" and len(val) > 100: e.append("grammar_note > 100 chars")
    if bank == "redemittel":
        if col == "redemittel_de":
            n = len([w for w in val.split() if w.strip(".…") != ""])
            if not 3 <= n <= 16: e.append(f"redemittel_de has {n} words (3-16)")
        if col == "grammar_note" and len(val) > 100: e.append("grammar_note > 100 chars")
    if col == "level" and bank == "vocab" and val not in LEVEL: pass
    return e


def clean(s):
    return re.sub(r"[^\w\s]", " ", s.lower().replace("ß", "ss")).split()


def warn_example(bank, r):
    """soft check: example contains target"""
    ex = (r.get("example_de") or "")
    if not ex: return None
    ex_t = clean(ex)
    if bank == "phrase":
        for seg in (r.get("phrase_de") or "").split("..."):
            seg = clean(seg)
            if seg and not any(ex_t[i:i + len(seg)] == seg for i in range(len(ex_t))):
                return "phrase not literally in example"
        return None
    if bank == "vocab":
        lem = (r.get("lemma") or "").lower().replace("ß", "ss")
        lem = re.sub(r"^sich ", "", lem)
        if not lem or " " in lem: return None
        stem = lem[: max(3, len(lem) - 3)] if r.get("pos") in ("verb", "adjective") else lem[: max(3, len(lem) - 2)]
        sep = (r.get("sep_prefix") or "").lower()
        text = " ".join(ex_t)
        if stem in text: return None
        if r.get("pos") == "verb":
            f = (r.get("pres_3sg") or "").split(" ")[0].lower()
            if f and f[: max(3, len(f) - 2)] in text: return None
        return "target not in example?"
    return None


# ---------------------------------------------------------------- rules block
def rules(bank):
    txt = README.read_text(encoding="utf-8")
    m = re.search(rf"### RULES:{RULE_KEY[bank]}\n(.*?)### END", txt, re.S)
    return m.group(1).strip() if m else "(rules block not found)"


# ---------------------------------------------------------------- commands
def need_arg(msg):
    raise SystemExit(msg)


def opt(name, default=None):
    if name in args:
        i = args.index(name); v = args[i + 1]; del args[i:i + 2]; return v
    return default


def flag(name):
    if name in args:
        args.remove(name); return True
    return False


def vocab_order(rows):
    k = lambda r: (0, int(r["study_order"]), r["id"]) if r.get("study_order") else (1, 0, r["id"])
    return sorted(rows, key=k)


def compact(bank, r, ref=None):
    parts = []
    for c in ORDER_COLS[bank]:
        v = r.get(c, "")
        if c == "freq_rank" and bank == "vocab" and not v: continue
        if v: parts.append(f"{c}={v}")
    if bank == "vocab" and not r.get("topic") and r.get("nw_chapter"):
        parts.append(f"nw_chapter={r['nw_chapter']}")
    line = " | ".join(parts)
    if bank == "vocab" and "ocr?" in r.get("check", "") and ref and r.get("nw_entry"):
        e = ref.get(r["nw_entry"])
        if e: line += f"\n   NWREF[{r['nw_entry']}] de_raw={e['de_raw']} | en_raw={e['en_raw']} | headword={e['headword']} | art={e['article']} | pl={e['plural_form']}"
    return line


def cmd_next():
    bank = args[0]; n = int(args[1]); rest = args[2:]
    del args[:2]
    a, b, wave, cat = opt("--from"), opt("--to"), opt("--wave"), opt("--cat")
    incomplete = flag("--incomplete")
    print(f"### RULES:{bank}\n{rules(bank)}\n### END\n")
    if bank == "redemittel":
        plan = load_plan(); _, rows = load_bank("redemittel")
        counts = {}
        for r in rows: counts[r["cat_id"]] = counts.get(r["cat_id"], 0) + 1
        for p in plan:
            if wave and p["wave"] != wave: continue
            if cat and p["cat_id"] != cat: continue
            have = counts.get(p["cat_id"], 0)
            if have < int(p["min_count"]) or cat:
                print(f"CATEGORY {p['cat_id']} wave={p['wave']} group={p['group']}\n category_de={p['category_de']} | category_en={p['category_en']}")
                print(f" cover: {p['cover']}\n exam_tasks: {p['exam_tasks']}\n level_range={p['level_range']} speaker(default)={p['speaker']} audience(default)={p['audience']}")
                print(f" min_count={p['min_count']} have={have} need_at_least={max(0, int(p['min_count']) - have)} (write up to {n} rows this step)")
                ex = [r for r in rows if r["cat_id"] == p["cat_id"]]
                print(f" already written ({len(ex)}):")
                for r in ex: print(f"   {r['id']} [{r['speaker']}/{r['function']}] {r['redemittel_de']}")
                return
        print("all categories complete"); return
    st = load_state(); done = set(st["reviewed"][bank])
    h, rows = load_bank(bank)
    if bank == "vocab":
        rows = vocab_order(rows)
        if a or b:
            lo, hi = int(a or 0), int(b or 10**9)
            rows = [r for r in rows if r["study_order"] and lo <= int(r["study_order"]) <= hi]
    sel = [r for r in rows if ((r["id"] in done and r["todo"]) if incomplete else r["id"] not in done)]
    ref = None
    if bank == "vocab":
        _, rr = read_csv(IN / "reference" / "netzwerk_entries.csv"); ref = {x["i"]: x for x in rr}
    print(f"{len(sel)} rows pending in selection; showing {min(n, len(sel))}")
    for r in sel[:n]: print(compact(bank, r, ref))


def cmd_apply():
    pf = Path(args[0]); patch = json.loads(pf.read_text(encoding="utf-8"))
    bank, cols, prow = patch["bank"], patch["cols"], patch["rows"]
    errs = []
    if bank not in FILES: raise SystemExit("unknown bank")
    if cols[0] != "id": errs.append("first col must be id")
    header, rows = load_bank(bank)
    unknown = [c for c in cols if c not in header]
    if unknown: errs.append(f"unknown columns {unknown}")
    locked = [c for c in cols if c in LOCKED[bank] and c != "id"]
    new_ok = bank == "redemittel"
    plan = {p["cat_id"]: p for p in load_plan()} if bank == "redemittel" else {}
    if bank == "redemittel":
        locked = [c for c in locked if c not in ("wave", "cat_id")]
        # wave/cat_id/id allowed only on creation (checked per row)
    elif locked: errs.append(f"locked columns {locked}")
    if errs: return fail(errs)
    byid = {r["id"]: r for r in rows}
    new_rows, changes, touched, planned = [], [], [], []
    next_n = max([int(r["id"][1:]) for r in rows] + [0]) + 1
    for k, pr in enumerate(prow):
        if len(pr) != len(cols): errs.append(f"row {k}: {len(pr)} values for {len(cols)} columns"); continue
        d = dict(zip(cols, pr)); rid = d["id"]
        if bank == "redemittel" and rid == "":
            if not {"wave", "cat_id", "category"} <= set(cols): errs.append(f"row {k}: new row needs wave, cat_id, category"); continue
            p = plan.get(d["cat_id"])
            if not p or p["wave"] != d["wave"] or p["category_de"] != d["category"]:
                errs.append(f"row {k}: wave/cat_id/category do not match plan"); continue
            cur = {c: "" for c in header}
            for c, v in d.items():
                if c != "id": cur[c] = "" if v == CLEAR else norm_enum(bank, c, v)
            cur["id"] = f"R{next_n:04d}"; next_n += 1
            cur["source"] = cur["source"] or "claude"
            for c in header:
                for m in cell_errors(bank, c, cur[c], cur["id"]): errs.append(f"{cur['id']}.{c}: {m}")
            cur["_new"] = True
            new_rows.append(cur); planned.append(cur); continue
        r = byid.get(rid)
        if not r: errs.append(f"row {k}: unknown id {rid!r}"); continue
        upd = dict(r)
        for c, v in d.items():
            if c == "id" or v == "": continue
            if bank == "redemittel" and c in ("wave", "cat_id") and v != r[c]:
                errs.append(f"{rid}.{c}: locked"); continue
            v2 = "" if v == CLEAR else norm_enum(bank, c, v)
            for m in cell_errors(bank, c, v2, rid): errs.append(f"{rid}.{c}: {m}")
            if v2 != r[c]:
                upd[c] = v2
                if r[c] != "": changes.append((bank, rid, c, r[c], v2))
        planned.append(upd); touched.append(rid)
    if errs: return fail(errs)
    # commit
    out_rows = []
    upd_by = {u["id"]: u for u in planned if not u.get("_new")}
    for r in rows: out_rows.append(upd_by.get(r["id"], r))
    for n in new_rows:
        n.pop("_new"); out_rows.append(n)
    for r in out_rows: r["todo"] = todo(bank, r) if "todo" in header else r.get("todo", "")
    write_csv_atomic(OUT / FILES[bank], header, out_rows)
    if changes:
        with open(OUT / "changes.csv", "a", encoding="utf-8", newline="") as f:
            w = csv.writer(f, lineterminator="\r\n")
            for c in changes: w.writerow(c)
    st = load_state(); st["patches"] += 1
    if bank in st["reviewed"]:
        s = set(st["reviewed"][bank]); s.update(touched); st["reviewed"][bank] = sorted(s)
    if st["patches"] % 20 == 0:
        bk = WORK / "backup" / datetime.datetime.now().strftime("%Y%m%d_%H%M%S"); bk.mkdir(parents=True, exist_ok=True)
        for f in FILES.values():
            if (OUT / f).exists(): shutil.copyfile(OUT / f, bk / f)
        for f in ("changes.csv", "progress.md"):
            if (OUT / f).exists(): shutil.copyfile(OUT / f, bk / f)
        shutil.copyfile(STATE, bk / "state.json") if STATE.exists() else None
    save_state(st)
    ids = touched + [n["id"] for n in new_rows]
    inc = [r["id"] for r in out_rows if r["id"] in set(ids) and r["todo"]]
    warns = []
    for r in out_rows:
        if r["id"] in set(ids):
            w = warn_example(bank, r)
            if w: warns.append(f"{r['id']}:{w}")
    print(f"applied {pf.name}: bank={bank} rows_touched={len(touched)} new_rows={len(new_rows)} cells_corrected={len(changes)} incomplete_in_patch={len(inc)} patches_total={st['patches']}")
    if inc: print("  incomplete:", ", ".join(f"{i}({next(r['todo'] for r in out_rows if r['id'] == i)})" for i in inc[:12]))
    if warns: print(f"  example warnings ({len(warns)}):", "; ".join(warns[:15]))


def fail(errs):
    print(f"REJECTED: {len(errs)} error(s), nothing written")
    for e in errs[:30]: print("  -", e)
    sys.exit(1)


def cmd_validate():
    final = flag("--final")
    errors, incomplete = [], {}
    for bank in ("vocab", "phrase", "redemittel"):
        p = OUT / FILES[bank]
        if not p.exists(): errors.append(f"{bank}: file missing"); continue
        header, rows = load_bank(bank)
        if bank == "redemittel" and header != RED_HEADER: errors.append("redemittel header differs")
        if bank in INPUTS:
            ih, irows = read_csv(IN / INPUTS[bank])
            if header != ih: errors.append(f"{bank}: header differs from input")
            if [r["id"] for r in rows] != [r["id"] for r in irows]: errors.append(f"{bank}: ids/order differ from input")
            else:
                for r, ir in zip(rows, irows):
                    for c in LOCKED[bank] - {"todo"}:
                        if r[c] != ir[c]: errors.append(f"{bank} {r['id']}.{c}: locked column changed")
        ids = [r["id"] for r in rows]
        if len(ids) != len(set(ids)): errors.append(f"{bank}: duplicate ids")
        inc = 0
        for r in rows:
            for c in header:
                for m in cell_errors(bank, c, r.get(c, ""), r["id"]):
                    if bank in INPUTS and c in LOCKED[bank]: continue
                    errors.append(f"{bank} {r['id']}.{c}: {m}")
            t = todo(bank, r)
            if (r.get("todo", "") if "todo" in header else "") != t and "todo" in header: errors.append(f"{bank} {r['id']}: todo stale")
            if t and not final: inc += 1
            elif t and final: errors.append(f"{bank} {r['id']}: incomplete ({t})")
        incomplete[bank] = inc
        seen = {}
        for r in rows:
            if bank == "vocab": key = (r["lemma"].lower(), r["pos"], r["article"])
            elif bank == "phrase": key = (r["phrase_de"].lower(),)
            else: key = (r["redemittel_de"].lower(),)
            if key in seen and not re.search(r"dup:|drop\?", r.get("check", "")) and (bank != "vocab" or True):
                errors.append(f"{bank} duplicate {r['id']} == {seen[key]}" + (" (flag dup:<id>)" if bank != "redemittel" else ""))
            else:
                seen.setdefault(key, r["id"])
        if bank == "redemittel":
            plan = load_plan(); cnt = {}
            for r in rows: cnt[r["cat_id"]] = cnt.get(r["cat_id"], 0) + 1
            pm = {p["cat_id"]: p for p in plan}
            for r in rows:
                p = pm.get(r["cat_id"])
                if not p or p["wave"] != r["wave"] or p["category_de"] != r["category"]: errors.append(f"redemittel {r['id']}: plan mismatch")
            if final:
                for p in plan:
                    if cnt.get(p["cat_id"], 0) < int(p["min_count"]): errors.append(f"redemittel {p['cat_id']}: {cnt.get(p['cat_id'], 0)} < {p['min_count']}")
                if rows:
                    wv = sum(1 for r in rows if r["variants"]) / len(rows)
                    if wv < 0.7: errors.append(f"redemittel: only {wv:.0%} rows have variants (<70%)")
                    core = sum(1 for r in rows if r["priority"] == "core") / len(rows)
                    if not 0.5 <= core <= 0.7: errors.append(f"redemittel: core share {core:.0%} (target ~60%)")
    print(f"validate{' --final' if final else ''}: errors={len(errors)} incomplete(not counted as errors)={incomplete}")
    for e in errors[:30]: print("  -", e)
    return len(errors)


def stats():
    res = {}
    for bank in ("vocab", "phrase", "redemittel"):
        if not (OUT / FILES[bank]).exists(): continue
        header, rows = load_bank(bank)
        empt = {}
        for r in rows:
            for c in (todo(bank, r).split("|") if todo(bank, r) else []): empt[c] = empt.get(c, 0) + 1
        res[bank] = dict(rows=len(rows), complete=sum(1 for r in rows if not todo(bank, r)), empty=empt,
                         flagged=sum(1 for r in rows if r.get("check")), rows_data=rows)
    return res


def cmd_status():
    write = flag("--write")
    st = load_state(); s = stats(); lines = []
    for bank, d in s.items():
        rev = len(st["reviewed"].get(bank, []))
        lines.append(f"- {bank}: rows={d['rows']} complete={d['complete']} reviewed={rev} flagged={d['flagged']} empty_cells={d['empty']}")
    plan = load_plan(); cnt = {}
    for r in s.get("redemittel", {}).get("rows_data", []): cnt[r["cat_id"]] = cnt.get(r["cat_id"], 0) + 1
    tot = sum(int(p["min_count"]) for p in plan)
    lines.append(f"- redemittel total {sum(cnt.values())}/{tot} min; categories done {sum(1 for p in plan if cnt.get(p['cat_id'], 0) >= int(p['min_count']))}/{len(plan)}")
    open_cat = [f"{p['cat_id']} {cnt.get(p['cat_id'], 0)}/{p['min_count']}" for p in plan if cnt.get(p['cat_id'], 0) < int(p['min_count'])]
    lines.append(f"- redemittel open categories: {len(open_cat)} (first: {', '.join(open_cat[:4])})")
    lines.append(f"- patches applied: {st['patches']}")
    print("\n".join(lines))
    if write:
        pm = OUT / "progress.md"
        old = pm.read_text(encoding="utf-8") if pm.exists() else ""
        m = re.search(r"<!--AUTO-->.*?<!--/AUTO-->", old, re.S)
        auto = "<!--AUTO-->\n## Status (auto, " + datetime.datetime.now().strftime("%Y-%m-%d %H:%M") + ")\n" + "\n".join(lines) + "\n<!--/AUTO-->"
        new = old.replace(m.group(0), auto) if m else auto + "\n\n" + old
        pm.write_text(new, encoding="utf-8")


def cmd_export():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    GROUP = {"vocab": "lane order_freq_first skip_reason kind level_basis nw_ref nw_chapter nw_entry goethe_level freq_rank freq_basis example_source todo".split(),
             "phrase": "order_hint words source level_hint goethe_level corpus_freq quality rater_note nw_ref old_bank_id todo".split(),
             "redemittel": []}
    NAMES = {"vocab": ("Vocabulary", "1F4E79"), "phrase": ("Phrases", "1E6B3A"), "redemittel": ("Redemittel", "7B3F00")}
    WRAP = {"example_de", "example_en", "meaning_en", "note", "collocations", "other_senses", "grammar_note", "variants", "redemittel_de", "exam_use"}
    data = {}
    for bank in ("vocab", "phrase", "redemittel"):
        header, rows = load_bank(bank); data[bank] = (header, rows)
        name, color = NAMES[bank]
        ws = wb.create_sheet(name)
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=color)
        for r in rows:
            ws.append([r.get(h, "") for h in header])
        for ci, h in enumerate(header, 1):
            L = get_column_letter(ci)
            w = 36 if h in WRAP else 14
            if h in ("id", "order_hint", "words", "study_order", "batch", "wave"): w = 10
            if h in ("lemma", "phrase_de", "function", "topic", "category", "nw_chapter"): w = 28
            ws.column_dimensions[L].width = w
            if h in GROUP[bank]:
                ws.column_dimensions[L].outlineLevel = 1; ws.column_dimensions[L].hidden = True
            for cell in ws[L][1:]:
                cell.number_format = "@" if h in ("id", "nw_ref") else "General"
                if h in WRAP: cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        if rows: ws.auto_filter.ref = ws.dimensions
    ws = wb.create_sheet("Summary")
    put = lambda *v: ws.append(list(v))
    put("German Bank - summary"); ws["A1"].font = Font(bold=True, size=14)
    put("Exported", datetime.date.today().isoformat()); put()
    put("Bank", "Rows", "Complete", "With check flags")
    for bank in ("vocab", "phrase", "redemittel"):
        h, rows = data[bank]
        put(NAMES[bank][0], len(rows), sum(1 for r in rows if not todo(bank, r)), sum(1 for r in rows if r.get("check")))
    put(); put("Vocabulary by level");
    vr = data["vocab"][1]
    for lv in LEVEL + [""]: put(lv or "(empty)", sum(1 for r in vr if r["level"] == lv))
    put(); put("Vocabulary by topic")
    for t in TOPICS + [""]: put(t or "(empty)", sum(1 for r in vr if r["topic"] == t))
    put(); put("Redemittel by wave / category", "Rows", "min_count")
    rr = data["redemittel"][1]
    for p in load_plan():
        put(f"W{p['wave']} {p['cat_id']} {p['category_de']}", sum(1 for r in rr if r["cat_id"] == p["cat_id"]), int(p["min_count"]))
    ws.column_dimensions["A"].width = 60
    wb.move_sheet("Summary", offset=0)
    tmp = OUT / "German_Bank.xlsx.tmp"
    wb.save(tmp); os.replace(tmp, OUT / "German_Bank.xlsx")
    print("exported German_Bank.xlsx:", {NAMES[b][0]: len(data[b][1]) for b in data})


def main():
    if not args: print(__doc__); return
    cmd = args.pop(0)
    if cmd == "setup": init_bank_final(); print("setup done")
    elif cmd == "status": cmd_status()
    elif cmd == "next": cmd_next()
    elif cmd == "apply": cmd_apply()
    elif cmd == "validate": sys.exit(1 if cmd_validate() else 0)
    elif cmd == "export": cmd_export()
    else: print(__doc__)


if __name__ == "__main__":
    main()
