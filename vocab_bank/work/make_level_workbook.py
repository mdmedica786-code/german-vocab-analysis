#!/usr/bin/env python3
"""make_level_workbook.py - builds STUDY_HERE/Word_Lists_by_Level.xlsx (one sheet per level) from the CSV banks.

  python3 work/make_level_workbook.py
"""
import csv
from collections import Counter
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
BANKS = ROOT / "bank_final"
OUT = ROOT.parent / "STUDY_HERE" / "Word_Lists_by_Level.xlsx"
LEVELS = ["A1", "A2", "B1", "B2", "C1"]
HEAD = PatternFill("solid", fgColor="2F6FED")


def rd(name):
    with open(BANKS / name, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def topic(t):
    return t.split(" ", 1)[1] if t and " " in t else t


def german(r):
    art = (r.get("article") or "").strip()
    return f"{art} {r['lemma']}" if r["pos"] == "noun" and art in ("der", "die", "das") else r["lemma"]


def forms(r):
    out = []
    if r["pos"] == "noun" and r.get("plural"):
        out.append("Pl.: " + r["plural"])
    if r["pos"] == "verb":
        f = ", ".join(x for x in (r.get("pres_3sg"), r.get("preterite"), r.get("perfect")) if x)
        if f:
            out.append(f)
        if r.get("gov"):
            out.append("+ " + r["gov"])
    if r.get("extra_forms"):
        out.append(r["extra_forms"])
    return " · ".join(out)


WORD_COLS = [("No.", 6), ("German", 28), ("Forms", 30), ("Type", 11), ("Meaning (English)", 34), ("Example (German)", 50),
             ("Example (English)", 46), ("Collocations", 40), ("Topic", 26), ("Medical", 8), ("Status", 10)]


def word_row(n, r):
    return [n, german(r), forms(r), r["pos"], r["meaning_en"], r.get("example_de", ""), r.get("example_en", ""),
            r.get("collocations", ""), topic(r["topic"]), "yes" if (r.get("medical") or r["topic"].startswith("M")) else "", ""]


def sheet(wb, title, cols, rows):
    ws = wb.create_sheet(title)
    ws.append([c for c, _ in cols])
    for row in rows:
        ws.append(row)
    for i, (_, w) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.zoomScale = 110
    return ws


def main():
    main_rows = rd("vocabulary_bank.csv")
    main_rows.sort(key=lambda r: int(r["study_order"]) if r.get("study_order", "").isdigit() else 10**9)
    b2 = rd("b2_vocabulary_bank.csv")
    b2.sort(key=lambda r: (r["topic"], r["lemma"].lower()))
    phr = rd("phrase_bank.csv")
    red = rd("redemittel_bank.csv")

    wb = Workbook()
    wb.remove(wb.active)
    study = [r for r in main_rows if not r.get("skip")]
    skipped = [r for r in main_rows if r.get("skip")]
    cnt = Counter(r["level"] for r in study)

    info = wb.create_sheet("Start here")
    lines = [["Word lists by level"], [""],
             ["One sheet per level. Use the filter arrows in the header row to filter by topic or type."],
             ["Type the status (new / learning / known) in the last column to track yourself."], [""],
             ["Sheet", "Words", "What it is"]]
    for lv in LEVELS:
        lines.append([lv, cnt[lv], "main bank, in study order" + (" (core B2 words)" if lv == "B2" else "")])
    lines += [["B2 plus", len(b2), f"{len(b2):,} extra B2 words for exams and daily life, no medical, none of them in the sheets above"],
              ["Phrases", len(phr), "useful phrases grouped by purpose"],
              ["Redemittel", len(red), "sentences for the speaking and writing exam"],
              ["Not in list", len(skipped), "words left out of the study list (country names, words you already know)"]]
    for ln in lines:
        info.append(ln)
    info["A1"].font = Font(bold=True, size=14)
    for c in info[6]:
        c.font = Font(bold=True)
    info.column_dimensions["A"].width = 16
    info.column_dimensions["B"].width = 10
    info.column_dimensions["C"].width = 90

    for lv in LEVELS:
        rows = [word_row(n, r) for n, r in enumerate([r for r in study if r["level"] == lv], 1)]
        sheet(wb, lv, WORD_COLS, rows)
    sheet(wb, "B2 plus", WORD_COLS, [word_row(n, r) for n, r in enumerate(b2, 1)])
    sheet(wb, "Phrases", [("No.", 6), ("Phrase", 38), ("Meaning (English)", 34), ("Level", 7), ("Purpose", 34), ("Example (German)", 50),
                          ("Example (English)", 46), ("Variants", 40), ("Grammar", 28)],
          [[n, r["phrase_de"], r["meaning_en"], r["level"], r["function"], r["example_de"], r["example_en"], r["variants"], r["grammar_note"]]
           for n, r in enumerate(sorted(phr, key=lambda r: (r["function"], r["level"])), 1)])
    sheet(wb, "Redemittel", [("No.", 6), ("Set", 34), ("Expression", 44), ("Meaning (English)", 36), ("Level", 7), ("Core", 7),
                             ("Example (German)", 50), ("Example (English)", 46), ("Variants", 40), ("Exam", 28)],
          [[n, f"{r['cat_id']} {r['category']}", r["redemittel_de"], r["meaning_en"], r["level"], "yes" if r["priority"] == "core" else "",
            r["example_de"], r["example_en"], r["variants"], r["exam_use"]] for n, r in enumerate(red, 1)])
    sheet(wb, "Not in list", WORD_COLS[:-1] + [("Reason", 40)],
          [word_row(n, r)[:-1] + [r.get("skip_reason", "")] for n, r in enumerate(skipped, 1)])
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(OUT, {lv: cnt[lv] for lv in LEVELS}, "b2 plus", len(b2))


if __name__ == "__main__":
    main()
