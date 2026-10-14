"""readme_sheet.py - adds a plain-English 'READ ME FIRST' sheet (what each column means) to the Excel exports."""
from openpyxl.styles import Alignment, Font, PatternFill

VOCAB = [
    ("lemma", "the word itself, in its dictionary form (verbs: infinitive; nouns: singular)"),
    ("article", "der / die / das for nouns (masculine / feminine / neuter)"),
    ("plural", "plural of the noun; (Pl.) = only used in the plural; – = no plural"),
    ("pos", "type of word: noun, verb, adjective, adverb, preposition ..."),
    ("pres_3sg / preterite / perfect", "verb forms: er/sie/es ...t  |  simple past  |  perfect (hat/ist + participle)"),
    ("sep_prefix", "the prefix that splits off in a sentence (ein- in 'einkaufen': ich kaufe ein)"),
    ("reflexive", "sich = the verb is used with sich (ich freue mich)"),
    ("gov", "what the word takes: case or preposition (e.g. auf + Akk.)"),
    ("meaning_en", "English meaning"),
    ("example_de / example_en", "one example sentence in German and its English translation"),
    ("collocations", "typical word combinations, separated by |"),
    ("level", "A1, A2, B1, B2 or C1"),
    ("topic", "T01-T16 = everyday topics, M1-M4 = medical topics"),
    ("medical", "core / related = medical vocabulary"),
    ("skip, skip_reason", "yes = NOT in the study list (country names, words you already know, duplicates ...)"),
    ("check", "a flag for me to look at: dup:B123 = same word as B123, drop? = Austrian/Swiss word, others = please verify"),
    ("exam_use", "where it helps: Lesen, Hören, Schreiben, Sprechen, Alltag ..."),
    ("freq_rank, study_order, batch, lane, kind, nw_*, goethe_level ...", "working columns (frequency, study order, source in Netzwerk/Goethe). Hidden in the main file; you can ignore them."),
]
OTHER = {
    "phrases": [("phrase_de", "the phrase"), ("function", "what it is for (agree, give reasons, ...)"), ("variants", "other ways to say it"), ("grammar_note", "short grammar hint")],
    "redemittel": [("category", "the speaking/writing situation"), ("redemittel_de", "the ready-made expression"), ("speaker", "Arzt / Patient / Pflege / Beide"), ("priority", "core = learn first")],
}
HOW = ["HOW TO USE THIS FILE",
       "1. Use the filter arrows in the first row (Data > Filter) to look at one level, one topic or one word type.",
       "2. To STUDY, do not use this file. Open STUDY_HERE/German_Study.html (flashcards) or WORD_LISTS/Word_Lists_by_Level.xlsx (clean lists, one sheet per level).",
       "3. This file is the complete data with every column. Do not sort and save over it: it is rebuilt by the tools."]


def add_readme_sheet(wb, kind="vocab"):
    ws = wb.create_sheet("READ ME FIRST", 0)
    ws.append([HOW[0]]); ws["A1"].font = Font(bold=True, size=14)
    for line in HOW[1:]:
        ws.append([line])
    ws.append([])
    ws.append(["Column", "What it means"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2F6FED")
    items = VOCAB + ([(k, v) for k, v in OTHER["phrases"]] if kind == "all" else [])
    if kind == "all":
        items += [(f"[Redemittel] {k}", v) for k, v in OTHER["redemittel"]]
    for k, v in items:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    return ws
